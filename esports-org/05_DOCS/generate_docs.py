#!/usr/bin/env python3
"""Erzeugt alle Vorlagen als Word-Dateien. Namen/Farben/Logo ändern -> Script neu starten.
Platzhalter in [eckigen Klammern] werden gelb markiert und müssen ausgefüllt werden."""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ORG = "VALIOUX"                      # <- Organisationsname (Arbeitsname)
LEGAL = "[Firmenname + Rechtsform, z. B. VALIOUX UG (haftungsbeschränkt)]"
ADDR = "[Straße, PLZ Ort]"
LOGO = os.path.join(os.path.dirname(__file__), "..", "01_BRAND", "logo-main.png")  # optional: wird automatisch genutzt
DARK = RGBColor(0x0A, 0x0B, 0x0E); SILVER = RGBColor(0x8E, 0x96, 0xA3); GREY = RGBColor(0x55, 0x5B, 0x66)
OUT = os.path.join(os.path.dirname(__file__), "docs")
os.makedirs(OUT, exist_ok=True)
DISCLAIMER = ("Mustervorlage – keine Rechtsberatung. Vor Verwendung durch eine Rechtsanwältin / einen Rechtsanwalt prüfen "
              "und an Land, Rechtsform und Einzelfall anpassen. Gelb markierte [Platzhalter] ausfüllen.")

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), hexcolor); tcPr.append(s)

def runs(p, text, bold=False, size=None, color=None, italic=False):
    for part in re.split(r'(\[[^\]]+\])', text):
        if not part: continue
        r = p.add_run(part); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        if part.startswith('[') and part.endswith(']'): r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def new_doc(title, subtitle=None):
    d = Document()
    sec = d.sections[0]; sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(2)
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri'); st.paragraph_format.space_after = Pt(5)
    # Kopfband (Logo + Name) – dunkel mit silberner Schrift
    t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells; shade(a, '0A0B0E'); shade(b, '0A0B0E')
    a.width = Cm(6); b.width = Cm(10.6)
    pa = a.paragraphs[0]
    if os.path.exists(LOGO): pa.add_run().add_picture(LOGO, height=Cm(1.5))
    else: runs(pa, "VLX", bold=True, size=24, color=RGBColor(0xE4, 0xEA, 0xF2))
    pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    runs(pb, ORG + "\n", bold=True, size=14, color=RGBColor(0xE4, 0xEA, 0xF2)); runs(pb, "ESPORTS ORGANISATION", size=8, color=SILVER)
    d.add_paragraph()
    p = d.add_paragraph(); runs(p, title.upper(), bold=True, size=18, color=DARK)
    if subtitle: runs(d.add_paragraph(), subtitle, size=10.5, color=GREY, italic=True)
    ln = d.add_paragraph(); ln.paragraph_format.space_after = Pt(8)
    pPr = ln._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', '8E96A3')): bt.set(qn('w:' + k), v)
    bd.append(bt); pPr.append(bd)
    # Fußzeile
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(fp, f"{ORG} · {LEGAL} · {ADDR}\n{DISCLAIMER}", size=7, color=SILVER)
    return d

def H(d, text): p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10); runs(p, text, bold=True, size=11.5, color=DARK); p.paragraph_format.keep_with_next = True
def P(d, text, indent=0): p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(indent); runs(p, text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; return p

def clauses(d, items, start=1):
    n = start
    for head, paras in items:
        H(d, f"§ {n} {head}"); n += 1
        for i, t in enumerate(paras if isinstance(paras, list) else [paras], 1):
            P(d, f"({i}) {t}" if isinstance(paras, list) and len(paras) > 1 else t)

def parties(d, a_label, b_label, a_txt=None, b_txt=None):
    t = d.add_table(rows=2, cols=2); t.style = 'Table Grid'
    for i, (lab, txt) in enumerate(((a_label, a_txt or f"{LEGAL}\n{ADDR}\nvertreten durch [Name, Funktion]"), (b_label, b_txt or "[Vor- und Nachname]\n[Straße, PLZ Ort]\n[Geburtsdatum] · Spielername/Handle: [Handle]"))):
        c0, c1 = t.rows[i].cells; shade(c0, 'E9ECF1'); runs(c0.paragraphs[0], lab, bold=True); runs(c1.paragraphs[0], txt)
        c0.width = Cm(4); c1.width = Cm(12.6)
    d.add_paragraph()

def signatures(d, a, b):
    d.add_paragraph(); t = d.add_table(rows=2, cols=2)
    for i, lab in enumerate((a, b)):
        runs(t.rows[0].cells[i].paragraphs[0], "\n\n____________________________\n[Ort, Datum]", color=GREY)
        runs(t.rows[1].cells[i].paragraphs[0], lab, bold=True)

def form_table(d, rows, header=None, widths=None):
    cols = len(rows[0]) if rows else len(header)
    t = d.add_table(rows=0, cols=cols); t.style = 'Table Grid'
    if header:
        r = t.add_row().cells
        for i, h in enumerate(header): shade(r[i], '0A0B0E'); runs(r[i].paragraphs[0], h, bold=True, color=RGBColor(0xE4, 0xEA, 0xF2), size=9.5)
    for row in rows:
        r = t.add_row().cells
        for i, v in enumerate(row):
            runs(r[i].paragraphs[0], v, size=9.5)
            if i == 0 and not header: shade(r[i], 'E9ECF1'); r[i].paragraphs[0].runs[0].bold = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph(); return t

def save(d, name): d.save(os.path.join(OUT, name + ".docx")); print("✓", name)

# ===================== VERTRÄGE =====================
def spielervertrag():
    d = new_doc("Spielervertrag (Pro)", "Professioneller Esports-Spieler · Titel: [Fortnite]")
    parties(d, "Organisation", "Spieler/in")
    P(d, f"zwischen {ORG} (nachfolgend „Organisation“) und dem/der Spieler/in (nachfolgend „Spieler“) wird Folgendes vereinbart:")
    clauses(d, [
     ("Gegenstand, Status", ["Der Spieler tritt für die Organisation als Mitglied des Pro-Teams im Titel [Fortnite] an und vertritt sie bei Turnieren, Scrims, Events und Medienterminen.",
        "Rechtsstatus: [angestellt / selbstständig]. Hinweis intern: Die Einordnung ist rechtlich zu prüfen (Weisungsbindung, Arbeitszeit, Scheinselbstständigkeit)."]),
     ("Laufzeit", ["Der Vertrag beginnt am [Datum] und läuft bis [Datum]. Probezeit: [x] Monate, Kündigung in der Probezeit mit [14] Tagen Frist.",
        "Verlängerungsoption der Organisation um [x] Monate zu unveränderten Konditionen bei Mitteilung bis [x] Wochen vor Ablauf."]),
     ("Pflichten des Spielers", ["Teilnahme an Training, Scrims und Turnieren nach Trainingsplan, mindestens [x] Stunden pro Woche.",
        "Pünktlichkeit, Fairplay, Befolgen der Anweisungen von Coach und Management im sportlichen Bereich.",
        "Einhaltung des Verhaltenskodex und der Social-Media-Richtlinie (Anlage 1).",
        "Meldung von Krankheit, Verhinderung und Reisen ohne schuldhaftes Zögern.",
        "Verzicht auf Match-Fixing, Wetten auf eigene Spiele, Cheating, Smurfing und jede Form unerlaubter Hilfsmittel."]),
     ("Vergütung", ["Monatliche Grundvergütung: [Betrag] € brutto/netto, zahlbar zum [letzten Werktag] des Monats auf das vom Spieler genannte Konto.",
        "Prämien: [x] % des Preisgelds der Organisation aus Turnieren, an denen der Spieler teilnahm; Verteilung nach Anlage 2.",
        "Leistungsbonus: [Betrag] € bei [Ziel, z. B. Qualifikation für XY]. Sachleistungen: [Equipment, Reisekosten, Unterkunft]."]),
     ("Equipment, Reisen, Spesen", ["Die Organisation stellt [Hardware/Peripherie] leihweise zur Verfügung; das Eigentum verbleibt bei der Organisation. Rückgabe bei Vertragsende in ordnungsgemäßem Zustand.",
        "Reisekosten zu genehmigten Terminen trägt die Organisation gegen Beleg nach der Spesenrichtlinie."]),
     ("Content, Stream, Sponsoring", ["Der Spieler streamt/produziert Inhalte im Umfang von mindestens [x] Stunden/Monat. Kanalrechte liegen beim Spieler; die Organisation erhält ein nicht ausschließliches Recht zur Nutzung von Inhalten mit Org-Bezug.",
        "Sponsoring-Verpflichtungen der Organisation (Logos, Produkte, Pflichttermine) sind vom Spieler zu beachten. Konkurrierende Sponsoren/Marken in Kategorien: [Liste].",
        "Der Spieler holt für eigene Sponsorings vorab Zustimmung ein, soweit Kategorien der Org-Partner betroffen sind."]),
     ("Bild-, Namens- und Persönlichkeitsrechte", "Der Spieler räumt der Organisation für die Vertragslaufzeit das Recht ein, Name, Handle, Bild und Stimme in Verbindung mit der Organisation für Marketing, Merchandise und Medien zu nutzen (Details in Anlage 3). Nachlaufend bleibt Archivnutzung bereits veröffentlichter Inhalte zulässig."),
     ("Vertraulichkeit", "Der Spieler wahrt Stillschweigen über interne Strategien, Verträge, Finanzen, Spielerdaten und Sponsorendetails, auch nach Vertragsende [x] Jahre."),
     ("Wettbewerbsverbot / Exklusivität", "Während der Laufzeit spielt der Spieler im Titel [Fortnite] ausschließlich für die Organisation. Wechselgespräche mit Dritten sind nur nach Vorabinformation der Organisation erlaubt, soweit rechtlich zulässig."),
     ("Transfer, Buyout", "Ein vorzeitiger Wechsel ist nur nach schriftlicher Freigabe möglich. Ablösesumme / Buyout: [Betrag] €. Die Organisation entscheidet innerhalb von [x] Tagen über Angebote."),
     ("Beendigung", ["Ordentliche Kündigung: [x] Monate zum Monatsende / ausgeschlossen bis [Datum].",
        "Außerordentliche Kündigung aus wichtigem Grund, insbesondere bei schweren Verstößen gegen Fairplay-Regeln, Verhaltenskodex oder Vertraulichkeit, bei Sperren durch Veranstalter/Epic Games und bei Rufschädigung.",
        "Kündigung bedarf der Textform. Equipment ist binnen [7] Tagen zurückzugeben; Zugänge (Accounts der Org) werden übergeben."]),
     ("Minderjährige", "Ist der Spieler minderjährig, ist die Zustimmung der gesetzlichen Vertreter (Anlage 4) Wirksamkeitsvoraussetzung. Jugendschutz- und Arbeitszeitvorschriften werden beachtet."),
     ("Schlussbestimmungen", "Änderungen bedürfen der Textform. Es gilt [deutsches] Recht. Gerichtsstand: [Ort]. Sollten Bestimmungen unwirksam sein, bleibt der Vertrag im Übrigen wirksam (salvatorische Klausel). Anlagen: 1 Verhaltenskodex & Social Media, 2 Prämienverteilung, 3 Bildrechte, 4 Erziehungsberechtigte."),
    ]); signatures(d, "Organisation", "Spieler"); save(d, "01_Spielervertrag_Pro")

def trialvertrag():
    d = new_doc("Trial-/Probevereinbarung", "Kurzvertrag für Probetraining und Tryout")
    parties(d, "Organisation", "Kandidat/in")
    clauses(d, [
     ("Zweck und Dauer", "Der Kandidat nimmt vom [Datum] bis [Datum] am Probetraining der Organisation im Titel [Fortnite] teil, um sportliche und menschliche Eignung zu prüfen."),
     ("Rechte und Pflichten", ["Teilnahme an vereinbarten Trainings/Scrims, Einhaltung des Verhaltenskodex.", "Es besteht kein Anspruch auf Abschluss eines Spielervertrags. Eine Vergütung erfolgt nur, wenn schriftlich vereinbart: [keine / Pauschale [Betrag] €]."]),
     ("Wettbewerb und Exklusivität", "Während der Trial-Phase werden keine Verträge mit konkurrierenden Organisationen im selben Titel geschlossen, ohne die Organisation vorab zu informieren."),
     ("Vertraulichkeit & Bildrechte", "Interne Inhalte bleiben vertraulich. Foto/Video aus dem Trial darf nur mit Zustimmung veröffentlicht werden."),
     ("Ende", "Die Vereinbarung endet automatisch mit Ablauf; beide Seiten können jederzeit ohne Angabe von Gründen beenden. Ergebnis wird innerhalb von [x] Tagen mitgeteilt."),
    ]); signatures(d, "Organisation", "Kandidat"); save(d, "02_Trial_Vereinbarung")

def staffvertrag():
    d = new_doc("Staff-Vertrag", "Coach · Analyst · Manager · Social/Content · Admin")
    parties(d, "Organisation", "Staff-Mitglied", b_txt="[Name]\n[Anschrift]\nPosition: [Head Coach / Analyst / Manager / Community / Editor]")
    clauses(d, [
     ("Position und Aufgaben", "Aufgabenbeschreibung: [Aufgaben]. Berichtslinie: [an Head of Operations/Gründer]."),
     ("Laufzeit / Probezeit", "Beginn [Datum]; befristet bis [Datum] / unbefristet. Probezeit [x] Monate."),
     ("Vergütung", "[Betrag] € monatlich brutto / Honorar [Betrag] € pro [Monat/Projekt]. Boni: [Regelung]. Aufwandserstattung nach Beleg."),
     ("Arbeitszeit, Ort", "Wochenarbeitszeit [x] Std., Einsatz [remote/Büro/Bootcamp], Erreichbarkeit [Zeiten]."),
     ("Geistiges Eigentum", "Alle im Rahmen der Tätigkeit geschaffenen Werke (Designs, Videos, Analysen, Texte, Code) stehen der Organisation zu; der Staff räumt alle erforderlichen Nutzungsrechte exklusiv und übertragbar ein."),
     ("Vertraulichkeit & Datenschutz", "Verpflichtung auf Geheimhaltung und Datenschutz (DSGVO). Zugang zu Org-Accounts nur dienstlich; Passwörter nur über den Passwortmanager der Org."),
     ("Kündigung", "Frist [x] Wochen zum Monatsende; außerordentliche Kündigung bei wichtigem Grund. Rückgabe von Geräten/Zugängen."),
    ]); signatures(d, "Organisation", "Staff-Mitglied"); save(d, "03_Staff_Vertrag")

def creatorvertrag():
    d = new_doc("Content-Creator-Vertrag", "Streamer · YouTuber · TikToker")
    parties(d, "Organisation", "Creator", b_txt="[Name]\n[Anschrift]\nKanäle: [Twitch/YouTube/TikTok/X-Links]")
    clauses(d, [
     ("Leistungen", "Mindestens [x] Streams/Videos/Posts je Monat, davon [x] Org-gebrandete Inhalte; Tragen des Org-Logos im Stream/Overlay; Nennung der Org in Titeln/Beschreibungen."),
     ("Vergütung", "Fix [Betrag] € monatlich + Umsatzbeteiligung [x] % auf [Merch/Codes/Affiliate]; Abrechnung monatlich bis zum [x]. Werktag."),
     ("Kanalrechte", "Kanäle und Kanalname bleiben Eigentum des Creators. Die Org erhält Nutzungsrecht an Org-bezogenen Inhalten für Marketing."),
     ("Werbung & Kennzeichnung", "Werbung wird nach geltendem Recht gekennzeichnet (z. B. „Werbung“/„Anzeige“). Konkurrenzkategorien: [Liste]; Vorabfreigabe für Sponsoren in den Bereichen [Liste]."),
     ("Verhalten", "Keine diskriminierenden, strafbaren oder rufschädigenden Inhalte; Verstöße berechtigen zur außerordentlichen Kündigung."),
     ("Laufzeit / Kündigung", "Laufzeit [x] Monate ab [Datum]; ordentliche Kündigung [x] Wochen zum Monatsende."),
    ]); signatures(d, "Organisation", "Creator"); save(d, "04_Creator_Vertrag")

def sponsoring():
    d = new_doc("Sponsoring-Vertrag", "Partnerschaft Organisation × Sponsor")
    parties(d, "Organisation", "Sponsor", b_txt="[Firma]\n[Anschrift]\nvertreten durch [Name]")
    clauses(d, [
     ("Gegenstand", "Der Sponsor unterstützt die Organisation als [Titelsponsor / Hauptpartner / Partner / Supplier] in den Titeln [Fortnite] und weiteren Aktivitäten."),
     ("Laufzeit", "Vom [Datum] bis [Datum]; Verlängerung nach gesonderter Vereinbarung. Vorrangige Verhandlungsoption für [x] Wochen."),
     ("Leistungen der Organisation", ["Logoplatzierung: Trikot [Position/Größe], Website, Overlay, Banner, Social [x] Posts/Monat.", "Content-Integration: [x] Videos/Streams, Rabattcode [CODE], Event-Auftritte [x].", "Reporting: monatlich Reichweiten- und Klickzahlen."]),
     ("Vergütung", "[Betrag] € [netto] zuzüglich USt, fällig [in Raten: Datum/Betrag]. Sachleistungen: [Produkte]. Zahlungsziel [14] Tage."),
     ("Kategorien-Exklusivität", "Kategorie: [z. B. Energy Drinks]. Die Org nimmt keine konkurrierenden Sponsoren in dieser Kategorie."),
     ("Marken- und Nutzungsrechte", "Gegenseitige, nicht übertragbare Lizenz zur Nutzung von Logos/Marken ausschließlich nach Vorgaben der Brand-Guidelines; Freigabe von Materialien innerhalb [x] Werktagen."),
     ("Compliance", "Beachtung von Werbe- und Jugendschutzrecht; Verzicht auf Sponsoring in Verbindung mit Glücksspiel/Alkohol ohne gesonderte Vereinbarung: [Regelung]."),
     ("Kündigung", "Außerordentlich bei Zahlungsverzug [x] Tage, schwerer Rufschädigung, Auflösung des Teams/Titels. Rückerstattung anteilig bei nicht erbrachten Leistungen."),
     ("Haftung", "Haftung nach den gesetzlichen Vorschriften; bei leichter Fahrlässigkeit nur für vertragswesentliche Pflichten und vorhersehbare Schäden."),
    ]); signatures(d, "Organisation", "Sponsor"); save(d, "05_Sponsoring_Vertrag")

def termsheet():
    d = new_doc("Investor Term Sheet", "Unverbindliche Eckpunkte (außer Vertraulichkeit & Exklusivität)")
    parties(d, "Gesellschaft", "Investor", b_txt="[Name / Gesellschaft]\n[Anschrift]\nvertreten durch [Name]")
    form_table(d, [
     ["Investitionssumme", "[Betrag] € in [Tranche 1: Betrag/Datum], [Tranche 2: Betrag/Meilenstein]"],
     ["Bewertung", "Pre-Money-Bewertung [Betrag] €; Beteiligung des Investors nach Finanzierung: [x] % (Cap-Table in Anlage)"],
     ["Instrument", "[Kapitalerhöhung (Geschäftsanteile) / Wandeldarlehen / stille Beteiligung / Revenue Share]"],
     ["Verwendung der Mittel", "Spielergehälter [x] %, Bootcamp/Reisen [x] %, Content & Marketing [x] %, Admin/Recht [x] %, Reserve [x] %"],
     ["Governance", "Beirat/Board: [x] Sitze Gründer, [x] Sitze Investor; zustimmungspflichtige Maßnahmen: [Budget > Betrag, Verkauf von Anteilen, Kreditaufnahme, Spieler-Transfers > Betrag]"],
     ["Informationsrechte", "Monatliches Reporting (Cashflow, Roster, Sponsoring-Pipeline), Jahresabschluss binnen [x] Monaten"],
     ["Verwässerungsschutz / Vorkaufsrecht", "[Pro-rata-Recht, Vorkaufs-/Mitverkaufsrecht (Tag-along), Drag-along bei [x] %]"],
     ["Exit / Rückzahlung", "[Exit-Erlös-Verteilung, Rückkaufoption der Gründer, Mindestlaufzeit [x] Jahre]"],
     ["Vesting Gründer", "[x] Jahre, Cliff [x] Monate"],
     ["Vertraulichkeit & Exklusivität", "Bindend: Vertraulichkeit; Exklusivität der Verhandlung [x] Wochen ab Unterschrift"],
     ["Bedingungen", "Due Diligence, Gesellschafterbeschluss, notarielle Beurkundung, Sponsoring-/Spielerverträge liegen vor"],
     ["Anwendbares Recht", "[deutsches Recht], Gerichtsstand [Ort]"],
    ]); signatures(d, "Gesellschaft", "Investor"); save(d, "06_Investor_TermSheet")

def investorvertrag():
    d = new_doc("Investorenvertrag / Beteiligungsvereinbarung", "Rahmenvertrag – Details notariell/anwaltlich ausgestalten")
    parties(d, "Gesellschaft", "Investor", b_txt="[Name / Gesellschaft]\n[Anschrift]")
    clauses(d, [
     ("Beteiligung", "Der Investor beteiligt sich mit [Betrag] € an der Gesellschaft und erhält [x] % der Geschäftsanteile / [Instrument]. Einzahlung binnen [x] Tagen nach Unterzeichnung/Beurkundung."),
     ("Mittelverwendung", "Die Mittel werden ausschließlich gemäß Investitionsplan (Anlage 1) verwendet. Abweichungen > [x] % bedürfen der Zustimmung des Investors."),
     ("Gewinn- und Erlösbeteiligung", "Ausschüttungen im Verhältnis der Beteiligung nach Bildung angemessener Rücklagen; Revenue-Share-Variante: [x] % der Netto-Einnahmen aus [Sponsoring/Merch/Preisgeld]."),
     ("Mitwirkung & Informationsrechte", "Beirat nach Term Sheet; Quartalsreporting; Einsicht in Bücher nach angemeldeter Prüfung."),
     ("Zustimmungspflichtige Geschäfte", "[Liste, z. B. Aufnahme von Krediten > Betrag, Veräußerung von Teams/Slots, Änderung des Gesellschaftszwecks]."),
     ("Wettbewerb, Interessenkonflikte", "Der Investor legt Beteiligungen an Wettbewerbern offen; Vertraulichkeit über Spielerverträge und Finanzen."),
     ("Übertragung von Anteilen", "Vorkaufsrecht der Gesellschaft/Gründer, Mitverkaufsrecht, Mitnahmepflicht nach Term Sheet."),
     ("Laufzeit, Exit, Kündigung", "Mindestlaufzeit [x] Jahre. Exit-Klauseln: [Verkauf, Rückkauf, Börsengang]. Ordentliche Kündigung ausgeschlossen, außerordentlich aus wichtigem Grund."),
     ("Garantien", "Die Gründer garantieren, dass Marken, Domains, Social-Media-Accounts und Verträge ordnungsgemäß bestehen und frei von Rechten Dritter sind."),
     ("Schlussbestimmungen", "Textform; salvatorische Klausel; [deutsches] Recht; Gerichtsstand [Ort]. Anlagen: 1 Investitionsplan, 2 Cap-Table, 3 Gesellschafterliste."),
    ]); signatures(d, "Gesellschaft / Gründer", "Investor"); save(d, "07_Investorenvertrag")

def nda():
    d = new_doc("Geheimhaltungsvereinbarung (NDA)", "Gegenseitig")
    parties(d, "Partei A", "Partei B", b_txt="[Name / Firma]\n[Anschrift]")
    clauses(d, [
     ("Vertrauliche Informationen", "Alle nicht öffentlichen Informationen, insbesondere Verträge, Finanzdaten, Strategien, Spielerdaten, Sponsoring-Konditionen, Pläne zu Branding und Roster, unabhängig von der Form."),
     ("Pflichten", "Nutzung nur zum Zweck [Zweck]; Weitergabe nur an Personen mit Need-to-know, die gleichwertig gebunden sind; angemessene Schutzmaßnahmen."),
     ("Ausnahmen", "Öffentlich bekannte, rechtmäßig von Dritten erhaltene oder eigenständig entwickelte Informationen; gesetzliche oder behördliche Offenlegungspflichten (mit Vorabinformation)."),
     ("Dauer", "Die Pflichten gelten ab Unterzeichnung für [3] Jahre über das Ende der Zusammenarbeit hinaus."),
     ("Vertragsstrafe", "Bei schuldhaftem Verstoß: [Betrag] € je Fall (angemessen, nach billigem Ermessen konkretisierbar); weitergehende Ansprüche bleiben unberührt."),
     ("Recht und Gericht", "[Deutsches] Recht; Gerichtsstand [Ort]."),
    ]); signatures(d, "Partei A", "Partei B"); save(d, "08_NDA")

def aufhebung():
    d = new_doc("Aufhebungs- / Transfervereinbarung", "Vertragsende, Wechsel, Buyout")
    parties(d, "Organisation", "Spieler / Staff")
    clauses(d, [
     ("Beendigung", "Der Vertrag vom [Datum] endet einvernehmlich zum [Datum]. Der Spieler wird [ab Datum] freigestellt."),
     ("Abfindung / Ablöse", "[keine / Betrag € von aufnehmender Organisation [Name], fällig am [Datum]]. Offene Vergütung bis Vertragsende: [Betrag]."),
     ("Rückgabe", "Equipment, Zugänge und Org-Accounts werden bis [Datum] zurückgegeben; Social-Accounts der Org werden an das Management übergeben."),
     ("Kommunikation", "Gemeinsame Ankündigung am [Datum]. Beide Seiten unterlassen abwertende Äußerungen."),
     ("Vertraulichkeit & Erledigung", "Vertraulichkeit bleibt bestehen. Mit Erfüllung sind gegenseitige Ansprüche abgegolten, ausgenommen [Vorbehalte]."),
    ]); signatures(d, "Organisation", "Spieler / Staff"); save(d, "09_Aufhebung_Transfer")

def verhaltenskodex():
    d = new_doc("Verhaltenskodex, Teamregeln & Social-Media-Richtlinie", "Anlage 1 zu allen Spieler-, Staff- und Creator-Verträgen")
    clauses(d, [
     ("Respekt & Fairplay", ["Wir behandeln Mitspieler, Gegner, Staff, Community und Partner respektvoll.", "Null Toleranz für Rassismus, Sexismus, Homophobie, Mobbing, Belästigung und Hassrede.", "Cheating, Boosting, Account-Sharing, Match-Fixing und Wetten auf eigene Matches sind verboten."]),
     ("Training & Verfügbarkeit", ["Pünktlichkeit bei Training, Scrims und Terminen; Absagen mindestens [x] Stunden vorher.", "Gesunder Lebensstil: Schlaf, Bewegung, Pausen; Unterstützung bei mentaler Gesundheit ist Priorität."]),
     ("Social Media", ["Posts, Streams und Kommentare spiegeln die Org wider – keine politischen Hetz-, Glücksspiel- oder Alkoholwerbung ohne Freigabe.", "Vor Sponsor-/Kooperationsposts Rücksprache mit dem Management. Werbung klar kennzeichnen.", "Keine Veröffentlichung interner Inhalte (Screenshots, Chats, Strategien, Verträge)."]),
     ("Streaming", "Stream-Sniping verhindern, Stream-Delay nach Turniervorgabe; Overlays mit Org-Branding; kein Streamen interner Calls."),
     ("Auftreten bei Events", "Teamkleidung und Branding, Alkoholverbot bei Terminen, respektvoller Umgang mit Fans und Presse."),
     ("Sanktionen", "Verwarnung → Abmahnung → Freistellung/Kündigung. Schwere Verstöße können sofort zur Kündigung führen. Meldung von Vorfällen vertraulich an [E-Mail/Ticket]."),
    ]); signatures(d, "Organisation", "Unterzeichner/in"); save(d, "10_Verhaltenskodex_SocialMedia")

def bildrechte():
    d = new_doc("Einwilligung Bild- & Persönlichkeitsrechte", "Anlage 3 zum Spieler-/Staff-/Creator-Vertrag")
    P(d, f"Ich willige ein, dass {ORG} Fotos, Videos, Audioaufnahmen, Name, Spielername/Handle und Spielstatistiken zu folgenden Zwecken nutzen darf: Website, Social Media, Streams, Presse, Merchandise, Sponsoring-Materialien, Ausstellungen/Events.")
    clauses(d, [
     ("Umfang", "Räumlich unbeschränkt, zeitlich für die Dauer der Zusammenarbeit; bereits veröffentlichte Inhalte dürfen als Archiv bestehen bleiben. Bearbeitung (Schnitt, Grafik) ist zulässig, entstellende Darstellungen nicht."),
     ("Merchandise", "Nutzung des Namens/Handles auf Merch nur nach Absprache; Beteiligung: [x] % des Netto-Erlöses mit Spielerbezug."),
     ("Widerruf", "Widerruf mit Wirkung für die Zukunft aus wichtigem Grund möglich; bereits gedruckte/produzierte Materialien müssen nicht zurückgerufen werden."),
     ("Datenschutz", "Ich wurde über Zwecke der Verarbeitung, Rechte (Auskunft, Löschung, Widerruf) und Ansprechpartner [Datenschutzkontakt] informiert."),
    ]); signatures(d, "Organisation", "Einwilligende/r"); save(d, "11_Einwilligung_Bildrechte")

def eltern():
    d = new_doc("Einwilligung der Erziehungsberechtigten", "Minderjährige Spieler/Talente · Anlage 4")
    parties(d, "Minderjährige/r", "Erziehungsberechtigte/r", a_txt="[Name] · geb. [Datum]\nSpielername: [Handle]", b_txt="[Name]\n[Anschrift]\n[Telefon/E-Mail]")
    clauses(d, [
     ("Einwilligung", f"Ich erkläre mich einverstanden, dass mein Kind für {ORG} spielt, trainiert und an Turnieren teilnimmt, und genehmige den Abschluss des Vertrags vom [Datum]."),
     ("Reisen & Events", "Reisen, Bootcamps und Events werden vorab mit mir abgestimmt. Begleitperson: [Name]. Notfallkontakt: [Telefon]."),
     ("Jugendschutz", "Trainingszeiten orientieren sich an Schul-/Ausbildungspflichten; Pausen werden eingehalten; keine Teilnahme an Veranstaltungen mit Alterskennzeichnung über dem Alter des Kindes."),
     ("Bild- und Social-Media-Rechte", "Ich stimme der Nutzung von Foto/Video/Name gemäß Anlage 3 zu [ja / nein, eingeschränkt auf: ___]."),
     ("Vergütung & Konto", "Zahlungen erfolgen auf [Konto des Erziehungsberechtigten]."),
    ]); signatures(d, "Organisation", "Erziehungsberechtigte/r"); save(d, "12_Einwilligung_Erziehungsberechtigte")

# ===================== FORMULARE =====================
def formular_bewerbung():
    d = new_doc("Bewerbungsformular", "Spieler · Academy · Staff · Creator")
    form_table(d, [["Name / Alter", "[ ]"], ["Spielername / Epic-Name", "[ ]"], ["Discord / X / Twitch / YouTube", "[ ]"], ["Land / Zeitzone / Sprachen", "[ ]"],
        ["Position / Rolle", "[ ]"], ["Titel / Modus", "[Fortnite – Zero Build / Build / Reload / Duos / Trio]"], ["Aktueller Rang / PR / FNCS-Ergebnisse", "[ ]"],
        ["Bisherige Teams", "[ ]"], ["Verfügbarkeit (Tage/Zeiten)", "[ ]"], ["Hardware / Internet (Ping, Mbit/s)", "[ ]"],
        ["Warum VALIOUX? Ziele?", "[ ]"], ["Link zu Highlights / Gameplay", "[ ]"], ["Bei Minderjährigen: Erziehungsberechtigte/r", "[ ]"]])
    P(d, "Mit dem Absenden willige ich in die Verarbeitung meiner Daten zum Zweck des Auswahlverfahrens ein. Löschung nach Abschluss des Verfahrens, spätestens nach [6] Monaten.")
    save(d, "F01_Bewerbung")

def formular_tryout():
    d = new_doc("Tryout-Bewertungsbogen", "Scoring 1–5 (5 = hervorragend)")
    form_table(d, [["Mechanik (Aim, Edit, Movement)", "[ ]", "[ ]"], ["Spielverständnis / Rotation", "[ ]", "[ ]"], ["Kommunikation & Callouts", "[ ]", "[ ]"], ["Teamfit / Attitude", "[ ]", "[ ]"],
        ["Coachability / Lernbereitschaft", "[ ]", "[ ]"], ["Konstanz & Mentale Stärke", "[ ]", "[ ]"], ["Verfügbarkeit / Zuverlässigkeit", "[ ]", "[ ]"], ["Content-/Reichweitenpotenzial", "[ ]", "[ ]"]],
        header=["Kriterium", "Score 1–5", "Notizen"], widths=[7, 3, 6.6])
    P(d, "Kandidat: [Name] · Datum: [ ] · Bewerter: [ ] · Gesamtscore: [ ] · Empfehlung: [Aufnahme / Trial verlängern / Absage]")
    save(d, "F02_Tryout_Bewertung")

def formular_spesen():
    d = new_doc("Spesen- / Reisekostenabrechnung")
    P(d, "Name: [ ] · Anlass: [ ] · Zeitraum: [ ] · Kostenstelle: [Pro / Academy / Content / Admin]")
    form_table(d, [["[ ]", "[ ]", "[ ]", "[ ] €", "[ ]"] for _ in range(8)], header=["Datum", "Beleg-Nr.", "Beschreibung", "Betrag", "Kategorie"], widths=[2.5, 2.5, 7, 2, 2.6])
    P(d, "Summe: [ ] € · Auszahlung auf IBAN: [ ] · Belege im Anhang · Genehmigt von: [ ] Datum: [ ]")
    save(d, "F03_Spesenabrechnung")

def formular_equipment():
    d = new_doc("Equipment-Ausgabe / Rückgabe")
    form_table(d, [["[PC/Monitor/Maus/Headset/Tastatur]", "[ ]", "[ ]", "[ ]", "[ ]"] for _ in range(6)], header=["Gerät", "Seriennr.", "Zustand", "Wert €", "Ausgabe / Rückgabe"], widths=[4.5, 3, 3, 2, 4.1])
    P(d, "Der Empfänger bestätigt, die Geräte pfleglich zu behandeln, nur dienstlich zu nutzen und bei Vertragsende zurückzugeben. Bei Verlust/Beschädigung durch Vorsatz oder grobe Fahrlässigkeit: Ersatz zum Zeitwert. Unterschrift: [ ] Datum: [ ]")
    save(d, "F04_Equipment")

def formular_stammdaten():
    d = new_doc("Stammdatenblatt (Spieler/Staff)")
    form_table(d, [[x, "[ ]"] for x in ["Vor-/Nachname", "Geburtsdatum / Staatsangehörigkeit", "Anschrift", "Telefon / E-Mail / Discord", "Notfallkontakt (Name/Telefon)", "IBAN / Kontoinhaber", "Steuer-ID / Sozialversicherung", "Konfektionsgröße Trikot / Hoodie", "Allergien / Hinweise (freiwillig)", "Reisepass-/Ausweisnummer (für Visa)", "Social Handles"]])
    P(d, "Daten werden ausschließlich für Vertragsdurchführung, Zahlungen und Organisation verarbeitet und nach Vertragsende fristgemäß gelöscht.")
    save(d, "F05_Stammdaten")

def formular_abwesenheit():
    d = new_doc("Abwesenheits-/Urlaubsantrag")
    form_table(d, [["Name / Team", "[ ]"], ["Zeitraum von – bis", "[ ]"], ["Grund (optional)", "[ ]"], ["Betroffene Termine / Turniere", "[ ]"], ["Vertretung / Ersatz", "[ ]"], ["Genehmigt (Coach / Management)", "[ ] Datum: [ ]"]])
    save(d, "F06_Abwesenheit")

def formular_vorfall():
    d = new_doc("Vorfallmeldung (vertraulich)", "Verhalten · Belästigung · Cheating-Verdacht · Sicherheitsvorfall")
    form_table(d, [["Datum / Uhrzeit", "[ ]"], ["Meldende Person (optional anonym)", "[ ]"], ["Betroffene Personen", "[ ]"], ["Beschreibung des Vorfalls", "[ ]"], ["Beweise (Screens, Links, Clips)", "[ ]"], ["Bereits ergriffene Maßnahmen", "[ ]"], ["Bearbeitung durch / Status", "[ ]"]])
    save(d, "F07_Vorfallmeldung")

def org_beschreibung():
    d = new_doc("Organisations-Profil", "Über uns · Pressetext · Sponsoren-Intro")
    clauses(d, [
     ("Kurzbeschreibung (160 Zeichen)", f"{ORG} – Fortnite-Esports-Organisation. Pro-Team, Academy, Creators und eigene Shows. Fearless. Elevated."),
     ("Mittellang", f"{ORG} ist eine Esports-Organisation mit Fokus auf Fortnite. Wir vereinen ein Pro-Team, eine Academy für kommende Talente und ein Creator-Netzwerk mit eigenen Live-Formaten. Unser Anspruch: Leistung, Haltung und Style – auf dem Server, im Stream und in der Community."),
     ("Lang / Pressetext", f"{ORG} wurde gegründet, um Spielern und Creators das Umfeld zu geben, das sie für die Weltspitze brauchen: professionelle Strukturen, klare Karrierepfade und eine Marke, die nach Qualität aussieht. Das Pro-Team tritt in den wichtigsten Fortnite-Wettbewerben an, die Academy entwickelt Talente systematisch in Richtung Pro-Roster, und Creator wie Shows machen {ORG} täglich live erlebbar – auf Twitch, YouTube, TikTok und X. Partner profitieren von einer jungen, aktiven Community und hochwertigem Content."),
     ("Mission", "Talente fördern, Titel gewinnen, Community begeistern."),
     ("Werte", "Fearless – wir gehen voran. Elevated – wir arbeiten auf Top-Niveau. Together – Team vor Ego. Authentic – echte Menschen, echte Stories."),
     ("Fakten-Box", "Gegründet: [Jahr] · Sitz: [Ort] · Titel: Fortnite · Rosters: Pro, Academy, Talents · Kontakt: [business@…] · Socials: [Links]"),
    ]); save(d, "00_Organisations_Profil")

if __name__ == "__main__":
    for f in (org_beschreibung, spielervertrag, trialvertrag, staffvertrag, creatorvertrag, sponsoring, termsheet, investorvertrag, nda, aufhebung, verhaltenskodex, bildrechte, eltern,
              formular_bewerbung, formular_tryout, formular_spesen, formular_equipment, formular_stammdaten, formular_abwesenheit, formular_vorfall): f()
