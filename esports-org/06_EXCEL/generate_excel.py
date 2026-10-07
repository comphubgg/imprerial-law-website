#!/usr/bin/env python3
"""Erzeugt die Management-Arbeitsmappe. ORG ändern -> neu starten."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from datetime import date

ORG = "VALIOUX"
BG = "0A0B0E"; SIL = "E4EAF2"; MID = "8E96A3"; ROW = "F3F5F8"
hf = Font(bold=True, color=SIL, name="Calibri", size=10); hfill = PatternFill("solid", fgColor=BG)
inp = PatternFill("solid", fgColor="FFFBE0"); calc = PatternFill("solid", fgColor=ROW)
thin = Side(style="thin", color="D5D9E0"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00 "€"'; PCT = '0.0%'; DT = 'DD.MM.YYYY'
wb = Workbook()

def sheet(name, title, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name; ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = "8E96A3"
    ws["A1"] = f"{ORG} · {title}"; ws["A1"].font = Font(bold=True, size=16, color=BG)
    ws["A2"] = "Gelb = Eingabe · Grau = Formel (nicht überschreiben)"; ws["A2"].font = Font(italic=True, size=9, color=MID)
    for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    return ws

def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c); x.font = hf; x.fill = hfill; x.border = box
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32; ws.freeze_panes = ws.cell(row=row + 1, column=3)

def style_rows(ws, r0, r1, ncol, fmts=None, formula_cols=()):
    for r in range(r0, r1 + 1):
        for c in range(1, ncol + 1):
            x = ws.cell(row=r, column=c); x.border = box
            x.fill = calc if c in formula_cols else inp
            if fmts and c in fmts: x.number_format = fmts[c]

def dv(ws, formula, rng):
    d = DataValidation(type="list", formula1=formula, allow_blank=True); ws.add_data_validation(d); d.add(rng)

# ---------------- Roster ----------------
R0, R1 = 5, 44
ro = sheet("Roster", "Roster & Gehälter", [6, 16, 20, 14, 12, 8, 7, 13, 13, 14, 13, 11, 15, 12, 13, 14, 13, 18], first=False)
header(ro, 4, ["#", "Handle", "Name", "Rolle", "Team", "Land", "Alter", "Vertrag ab", "Vertrag bis", "Grundgehalt / Monat", "Bonus / Jahr", "Preisgeld-Anteil", "Kosten / Jahr", "Tage bis Ablauf", "Status", "Kündigungsfrist", "Buyout €", "Discord / X"])
style_rows(ro, R0, R1, 18, {8: DT, 9: DT, 10: EUR, 11: EUR, 12: PCT, 13: EUR, 17: EUR}, formula_cols=(1, 13, 14, 15))
for r in range(R0, R1 + 1):
    ro.cell(r, 1, f"=IF(B{r}=\"\",\"\",ROW()-{R0-1})")
    ro.cell(r, 13, f"=IF(B{r}=\"\",\"\",J{r}*12+K{r})")
    ro.cell(r, 14, f"=IF(I{r}=\"\",\"\",I{r}-TODAY())")
    ro.cell(r, 15, f"=IF(I{r}=\"\",\"\",IF(N{r}<0,\"abgelaufen\",IF(N{r}<90,\"läuft aus\",\"aktiv\")))")
dv(ro, '"Pro,Academy,Talent,Staff,Creator"', f"E{R0}:E{R1}")
dv(ro, '"IGL,Fragger,Support,Flex,Sub,Coach,Analyst,Manager,Creator"', f"D{R0}:D{R1}")
ro.conditional_formatting.add(f"O{R0}:O{R1}", CellIsRule(operator="equal", formula=['"abgelaufen"'], fill=PatternFill("solid", bgColor="F8B4B4")))
ro.conditional_formatting.add(f"O{R0}:O{R1}", CellIsRule(operator="equal", formula=['"läuft aus"'], fill=PatternFill("solid", bgColor="FDE68A")))
ro.conditional_formatting.add(f"O{R0}:O{R1}", CellIsRule(operator="equal", formula=['"aktiv"'], fill=PatternFill("solid", bgColor="BBF7D0")))
for i, row in enumerate([("BeispielPro1", "Max Mustermann", "Fragger", "Pro", "DE", 19, date(2026, 1, 1), date(2026, 12, 31), 2500, 3000, 0.25, "3 Monate", 15000, "@bsp1"),
                         ("BeispielPro2", "Erika Beispiel", "IGL", "Pro", "DE", 21, date(2026, 1, 1), date(2027, 6, 30), 3000, 5000, 0.30, "3 Monate", 25000, "@bsp2"),
                         ("BeispielAcad", "Tim Talent", "Talent", "Academy", "DE", 16, date(2026, 6, 1), date(2026, 12, 31), 400, 0, 0.10, "1 Monat", 0, "@bsp3")]):
    cols = [2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 16, 17, 18]
    for c, v in zip(cols, row): ro.cell(R0 + i, c, v)
ro.cell(R1 + 2, 9, "Summe").font = Font(bold=True)
for c in (10, 11, 13): ro.cell(R1 + 2, c, f"=SUM({L(c)}{R0}:{L(c)}{R1})").number_format = EUR; ro.cell(R1 + 2, c).font = Font(bold=True)

# ---------------- Gehaltsrechner ----------------
gr = sheet("Preisgeld-Split", "Preisgeld- & Bonus-Rechner", [28, 16, 16, 16, 16])
gr["A4"] = "Turnier-Preisgeld gesamt (Team)"; gr["B4"] = 0; gr["B4"].number_format = EUR
gr["A5"] = "Anteil Organisation"; gr["B5"] = 0.30; gr["B5"].number_format = PCT
gr["A6"] = "Anteil Spieler (gesamt)"; gr["B6"] = "=1-B5"; gr["B6"].number_format = PCT
gr["A7"] = "Betrag Organisation"; gr["B7"] = "=B4*B5"; gr["B7"].number_format = EUR
gr["A8"] = "Betrag Spieler (gesamt)"; gr["B8"] = "=B4*B6"; gr["B8"].number_format = EUR
for r in (4, 5): gr.cell(r, 2).fill = inp
for r in (6, 7, 8): gr.cell(r, 2).fill = calc
header(gr, 10, ["Spieler", "Gewichtung", "Anteil %", "Auszahlung brutto", "Notiz"]); gr.freeze_panes = None
for r in range(11, 19):
    for c in range(1, 6): gr.cell(r, c).border = box; gr.cell(r, c).fill = inp if c in (1, 2, 5) else calc
    gr.cell(r, 3, f'=IF(B{r}="","",B{r}/SUM($B$11:$B$18))').number_format = PCT
    gr.cell(r, 4, f'=IF(B{r}="","",C{r}*$B$8)').number_format = EUR
gr["A20"] = "Hinweis: Gewichtung z. B. 1 = gleichmäßig. Steuern/Sozialabgaben separat prüfen (Steuerberater)."; gr["A20"].font = Font(italic=True, size=9, color=MID)

# ---------------- Budget ----------------
bu = sheet("Budget & Cashflow", "Jahresbudget & Cashflow", [30] + [12] * 12 + [14])
months = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
bu["A3"] = "Startkapital (Kontostand)"; bu["B3"] = 0; bu["B3"].number_format = EUR; bu["B3"].fill = inp
header(bu, 5, ["Kategorie"] + months + ["Summe"]); bu.freeze_panes = "B6"
inc = ["Sponsoring", "Preisgelder", "Merch / Shop", "Content / Werbeeinnahmen", "Investorengelder", "Sonstiges"]
exp = ["Spielergehälter (aus Roster)", "Staff / Coaches", "Bootcamp & Reisen", "Equipment & Hardware", "Marketing & Content-Produktion", "Software & Tools", "Recht, Steuer, Versicherung", "Turniergebühren / Slots", "Sonstiges"]
r = 6; bu.cell(r, 1, "EINNAHMEN").font = Font(bold=True); r += 1; inc0 = r
for n in inc:
    bu.cell(r, 1, n)
    for c in range(2, 14): bu.cell(r, c).fill = inp; bu.cell(r, c).number_format = EUR
    bu.cell(r, 14, f"=SUM(B{r}:M{r})"); r += 1
inc1 = r - 1; bu.cell(r, 1, "Summe Einnahmen").font = Font(bold=True); sumi = r
for c in range(2, 15): bu.cell(r, c, f"=SUM({L(c)}{inc0}:{L(c)}{inc1})")
r += 2; bu.cell(r, 1, "AUSGABEN").font = Font(bold=True); r += 1; exp0 = r
for k, n in enumerate(exp):
    bu.cell(r, 1, n)
    for c in range(2, 14):
        bu.cell(r, c).number_format = EUR
        if k == 0: bu.cell(r, c, f"=Roster!$J${R1+2}"); bu.cell(r, c).fill = calc
        else: bu.cell(r, c).fill = inp
    bu.cell(r, 14, f"=SUM(B{r}:M{r})"); r += 1
exp1 = r - 1; bu.cell(r, 1, "Summe Ausgaben").font = Font(bold=True); sume = r
for c in range(2, 15): bu.cell(r, c, f"=SUM({L(c)}{exp0}:{L(c)}{exp1})")
r += 2; bu.cell(r, 1, "Ergebnis (Monat)").font = Font(bold=True); res = r
for c in range(2, 15): bu.cell(r, c, f"={L(c)}{sumi}-{L(c)}{sume}")
r += 1; bu.cell(r, 1, "Kontostand kumuliert").font = Font(bold=True); kum = r
for c in range(2, 14): bu.cell(r, c, f"=$B$3+SUM($B${res}:{L(c)}{res})")
for rr in range(6, r + 1):
    for c in range(2, 15):
        x = bu.cell(rr, c); x.border = box; x.number_format = EUR
        if rr in (sumi, sume, res, kum): x.font = Font(bold=True); x.fill = calc
bu.conditional_formatting.add(f"B{kum}:M{kum}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))
bu.conditional_formatting.add(f"B{res}:N{res}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

# ---------------- Gründungskosten ----------------
gk = sheet("Gründungskosten", "Einmalige Gründungskosten", [38, 16, 16, 14, 40])
header(gk, 4, ["Posten", "Geplant €", "Ist €", "Status", "Hinweis"]); gk.freeze_panes = None
items = [("Rechtsform-Gründung (Notar, Handelsregister)", "Je nach Rechtsform/Land; Beratung einholen"), ("Rechtsberatung Verträge (Spieler, Sponsor, Investor)", "Vorlagen prüfen lassen"), ("Markenanmeldung (Wort-/Bildmarke)", "Klassen 9, 25, 35, 41; Recherche vorab"),
         ("Domain(s) & Hosting", "Hauptdomain + Schutzdomains"), ("Logo/Branding/Design (extern)", "Falls nicht selbst erstellt"), ("Trikot-Erstausstattung", "Design, Muster, Erstserie, Größen"), ("Hardware-Erstausstattung", "PCs, Monitore, Peripherie"),
         ("Versicherungen (Haftpflicht, Equipment, Unfall)", "Angebote vergleichen"), ("Steuerberater / Buchhaltung Setup", ""), ("Bankkonto, Zahlungsanbieter", ""), ("Software (Notion, Discord Bots, Adobe, OBS-Tools)", ""), ("Website Betrieb / Tools", ""), ("Launch-Marketing / Content-Paket", ""), ("Reserve (10–20 %)", "")]
for i, (n, h) in enumerate(items):
    r = 5 + i; gk.cell(r, 1, n); gk.cell(r, 5, h)
    for c in range(1, 6): gk.cell(r, c).border = box; gk.cell(r, c).fill = inp
    gk.cell(r, 2).number_format = EUR; gk.cell(r, 3).number_format = EUR
dv(gk, '"offen,beauftragt,bezahlt"', f"D5:D{4+len(items)}")
e = 5 + len(items); gk.cell(e + 1, 1, "Summe").font = Font(bold=True)
for c in (2, 3): gk.cell(e + 1, c, f"=SUM({L(c)}5:{L(c)}{e-1})").number_format = EUR

# ---------------- Sponsoren-Tracker ----------------
sp = sheet("Sponsoren", "Sponsoren-Pipeline", [22, 20, 18, 16, 14, 14, 14, 14, 36])
header(sp, 4, ["Firma", "Ansprechpartner", "Kontakt", "Phase", "Deal-Wert €", "Wahrscheinlichkeit", "Gewichtet €", "Nächster Schritt am", "Notiz"])
style_rows(sp, 5, 44, 9, {5: EUR, 6: PCT, 7: EUR, 8: DT}, formula_cols=(7,))
for r in range(5, 45): sp.cell(r, 7, f'=IF(E{r}="","",E{r}*F{r})')
dv(sp, '"Lead,Kontaktiert,Gespräch,Angebot,Verhandlung,Zugesagt,Abgelehnt"', "D5:D44")
sp["A46"] = "Pipeline gesamt"; sp["E46"] = "=SUM(E5:E44)"; sp["G46"] = "=SUM(G5:G44)"
sp["A47"] = "Zugesagt"; sp["E47"] = '=SUMIF(D5:D44,"Zugesagt",E5:E44)'
for c in ("E46", "G46", "E47"): sp[c].number_format = EUR; sp[c].font = Font(bold=True)

# ---------------- Investoren ----------------
iv = sheet("Investoren", "Investoren-Tracker", [22, 20, 18, 16, 14, 12, 14, 14, 34])
header(iv, 4, ["Investor", "Ansprechpartner", "Kontakt", "Phase", "Summe €", "Anteil %", "Pre-Money €", "Nächster Schritt am", "Notiz"])
style_rows(iv, 5, 24, 9, {5: EUR, 6: PCT, 7: EUR, 8: DT})
dv(iv, '"Lead,Pitch,Due Diligence,Term Sheet,Vertrag,Eingezahlt,Abgelehnt"', "D5:D24")
iv["A26"] = "Eingezahlt gesamt"; iv["E26"] = '=SUMIF(D5:D24,"Eingezahlt",E5:E24)'; iv["E26"].number_format = EUR; iv["E26"].font = Font(bold=True)
iv["A27"] = "Verbleibende Anteile Gründer"; iv["F27"] = '=1-SUMIF(D5:D24,"Eingezahlt",F5:F24)'; iv["F27"].number_format = PCT

# ---------------- Vertrags-Tracker ----------------
vt = sheet("Verträge", "Vertrags-Tracker mit Ablauf-Warnung", [10, 24, 22, 14, 14, 16, 14, 14, 28])
header(vt, 4, ["Nr.", "Partner", "Typ", "Beginn", "Ende", "Tage bis Ablauf", "Warnung", "Wert €", "Dokument-Link / Notiz"])
style_rows(vt, 5, 54, 9, {4: DT, 5: DT, 8: EUR}, formula_cols=(6, 7))
for r in range(5, 55):
    vt.cell(r, 6, f'=IF(E{r}="","",E{r}-TODAY())'); vt.cell(r, 7, f'=IF(E{r}="","",IF(F{r}<0,"ABGELAUFEN",IF(F{r}<=90,"bald fällig","ok")))')
dv(vt, '"Spielervertrag,Staff,Creator,Sponsoring,Investor,NDA,Sonstiges"', "C5:C54")
vt.conditional_formatting.add("G5:G54", CellIsRule(operator="equal", formula=['"ABGELAUFEN"'], fill=PatternFill("solid", bgColor="F8B4B4")))
vt.conditional_formatting.add("G5:G54", CellIsRule(operator="equal", formula=['"bald fällig"'], fill=PatternFill("solid", bgColor="FDE68A")))

# ---------------- Turniere ----------------
tu = sheet("Turniere & Preisgeld", "Turnier- & Preisgeld-Tracker", [14, 30, 14, 14, 14, 14, 14, 30])
header(tu, 4, ["Datum", "Turnier", "Team", "Platzierung", "Preisgeld €", "Org-Anteil %", "Org-Betrag €", "Notiz"])
style_rows(tu, 5, 64, 8, {1: DT, 5: EUR, 6: PCT, 7: EUR}, formula_cols=(7,))
for r in range(5, 65): tu.cell(r, 7, f'=IF(E{r}="","",E{r}*F{r})')
tu["D66"] = "Summe"; tu["E66"] = "=SUM(E5:E64)"; tu["G66"] = "=SUM(G5:G64)"
for c in ("E66", "G66"): tu[c].number_format = EUR; tu[c].font = Font(bold=True)

# ---------------- Equipment ----------------
eq = sheet("Equipment", "Equipment-Inventar", [10, 22, 22, 18, 14, 14, 18, 14, 22])
header(eq, 4, ["Nr.", "Gerät", "Hersteller/Modell", "Seriennummer", "Kaufdatum", "Wert €", "Ausgegeben an", "Zustand", "Rückgabe bis"])
style_rows(eq, 5, 64, 9, {5: DT, 6: EUR, 9: DT})
dv(eq, '"neu,gut,gebraucht,defekt,verloren"', "H5:H64")

# ---------------- Tryouts ----------------
tr = sheet("Tryouts", "Scouting & Tryout-Board", [20, 8, 14, 12, 12, 12, 12, 12, 12, 12, 12, 14, 26])
header(tr, 4, ["Kandidat", "Alter", "Kontakt", "Mechanik", "Game Sense", "Komm.", "Teamfit", "Coachability", "Konstanz", "Content", "Ø Score", "Status", "Notiz"])
style_rows(tr, 5, 54, 13, formula_cols=(11,))
for r in range(5, 55): tr.cell(r, 11, f'=IF(COUNT(D{r}:J{r})=0,"",ROUND(AVERAGE(D{r}:J{r}),2))')
dv(tr, '"Sichtung,Tryout,Trial,Angebot,Aufgenommen,Absage"', "L5:L54")
dv(tr, '"1,2,3,4,5"', "D5:J54")
tr.conditional_formatting.add("K5:K54", CellIsRule(operator="greaterThanOrEqual", formula=["4"], fill=PatternFill("solid", bgColor="BBF7D0")))

# ---------------- Content-Kalender ----------------
ck = sheet("Content-Kalender", "Content- & Post-Kalender", [14, 12, 16, 34, 18, 16, 16])
header(ck, 4, ["Datum", "Zeit", "Kanal", "Inhalt / Idee", "Format", "Verantwortlich", "Status"])
style_rows(ck, 5, 124, 7, {1: DT})
dv(ck, '"X,Twitch,YouTube,TikTok,Instagram,Discord,Website"', "C5:C124"); dv(ck, '"Idee,In Arbeit,Freigabe,Geplant,Veröffentlicht"', "G5:G124")

# ---------------- Reisekosten ----------------
rk = sheet("Reisekosten", "Reise- & Bootcamp-Kosten", [14, 26, 20, 12, 12, 12, 12, 14])
header(rk, 4, ["Datum", "Event / Ziel", "Person", "Reise €", "Hotel €", "Verpflegung €", "Sonstiges €", "Summe €"])
style_rows(rk, 5, 54, 8, {1: DT, 4: EUR, 5: EUR, 6: EUR, 7: EUR, 8: EUR}, formula_cols=(8,))
for r in range(5, 55): rk.cell(r, 8, f'=IF(COUNT(D{r}:G{r})=0,"",SUM(D{r}:G{r}))')
rk["G56"] = "Gesamt"; rk["H56"] = "=SUM(H5:H54)"; rk["H56"].number_format = EUR; rk["H56"].font = Font(bold=True)

# ---------------- Dashboard (zuerst) ----------------
da = wb.create_sheet("Dashboard", 0); da.sheet_view.showGridLines = False; da.sheet_properties.tabColor = BG
da.column_dimensions["A"].width = 4; da.column_dimensions["B"].width = 34; da.column_dimensions["C"].width = 22; da.column_dimensions["D"].width = 4; da.column_dimensions["E"].width = 34; da.column_dimensions["F"].width = 22
da.merge_cells("A1:F1"); da["A1"] = f"{ORG} · MANAGEMENT-DASHBOARD"; da["A1"].font = Font(bold=True, size=20, color=SIL); da["A1"].fill = hfill; da.row_dimensions[1].height = 42
da["A1"].alignment = Alignment(vertical="center", indent=1)
kp = [("B", "C", 3, "Spieler & Staff im Roster", f'=COUNTA(Roster!B{R0}:B{R1})', "0"), ("B", "C", 4, "Gehaltskosten / Monat", f"=Roster!J{R1+2}", EUR), ("B", "C", 5, "Personalkosten / Jahr", f"=Roster!M{R1+2}", EUR),
      ("B", "C", 6, "Verträge laufen aus (< 90 Tage)", '=COUNTIF(Roster!O5:O44,"läuft aus")', "0"), ("B", "C", 7, "Abgelaufene Verträge", '=COUNTIF(Roster!O5:O44,"abgelaufen")', "0"),
      ("E", "F", 3, "Sponsoring zugesagt", "=Sponsoren!E47", EUR), ("E", "F", 4, "Sponsoren-Pipeline (gewichtet)", "=Sponsoren!G46", EUR), ("E", "F", 5, "Investorengelder eingezahlt", "=Investoren!E26", EUR),
      ("E", "F", 6, "Preisgelder Org-Anteil", "='Turniere & Preisgeld'!G66", EUR), ("E", "F", 7, "Jahresergebnis (Budget)", f"='Budget & Cashflow'!N{res}", EUR)]
for a, b, r, lab, f, fmt in kp:
    da[f"{a}{r+1}"] = lab; da[f"{b}{r+1}"] = f; da[f"{b}{r+1}"].number_format = fmt
    da[f"{a}{r+1}"].font = Font(color=MID, size=10); da[f"{b}{r+1}"].font = Font(bold=True, size=14); da[f"{b}{r+1}"].alignment = Alignment(horizontal="right")
    for cc in (a, b): da[f"{cc}{r+1}"].border = Border(bottom=thin)
import os
_logo = os.path.join(os.path.dirname(__file__), "..", "01_BRAND", "logo-main.png")
if os.path.exists(_logo):
    from openpyxl.drawing.image import Image as XImage
    _im = XImage(_logo); _im.height = 52; _im.width = int(52 * 1248 / 649); da.add_image(_im, "E1")
da["B13"] = "Tabs: Roster · Preisgeld-Split · Budget & Cashflow · Gründungskosten · Sponsoren · Investoren · Verträge · Turniere · Equipment · Tryouts · Content-Kalender · Reisekosten"
da["B13"].font = Font(italic=True, size=9, color=MID)

del wb["Sheet"]
wb.save("VALIOUX_Management.xlsx"); print("OK", wb.sheetnames)
